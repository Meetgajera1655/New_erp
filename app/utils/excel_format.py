import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.filters import AutoFilter
from openpyxl import load_workbook


def format_excel(path, sheet_titles=None):
    """
    Apply professional formatting to all sheets in the Excel file at `path`.
    Optionally, provide a dict of {sheet_name: title_row_text} for custom titles.
    """
    wb = load_workbook(path)
    thin = Side(border_style="thin", color="CCCCCC")
    header_fill = PatternFill("solid", fgColor="D9E1F2")  # Light blue
    alt_fill = PatternFill("solid", fgColor="F2F2F2")    # Light gray
    for ws in wb.worksheets:
        # Title row
        title = sheet_titles[ws.title] if sheet_titles and ws.title in sheet_titles else ws.title
        ws.insert_rows(1)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
        cell = ws.cell(row=1, column=1)
        cell.value = title
        cell.font = Font(size=15, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        # Header row
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=2, column=col)
            cell.font = Font(bold=True, size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = header_fill
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
        # Freeze header
        ws.freeze_panes = ws[3][0]
        # Alternate row colors & borders
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
                if cell.row % 2 == 1:
                    cell.fill = alt_fill
                # Null values as "-"
                if cell.value is None or (isinstance(cell.value, float) and pd.isna(cell.value)):
                    cell.value = "-"
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    val = str(cell.value)
                except Exception:
                    val = ""
                if val is not None:
                    max_length = max(max_length, len(val))
            ws.column_dimensions[col_letter].width = max(12, min(max_length + 2, 50))
        # Date/number formatting
        for col in ws.iter_cols(min_row=3, max_row=ws.max_row):
            for cell in col:
                if isinstance(cell.value, str):
                    try:
                        pd.to_datetime(cell.value)
                        cell.number_format = "yyyy-mm-dd"
                    except Exception:
                        pass
                elif isinstance(cell.value, (int, float)):
                    if "amount" in (ws.cell(row=2, column=cell.column).value or "").lower() or "total" in (ws.cell(row=2, column=cell.column).value or "").lower():
                        cell.number_format = '#,##0.00'
        # Enable filter on header
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
